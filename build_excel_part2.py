import pandas as pd
import json
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

DATA_DIR = "/home/claude/pharma_project/data"
OUT_DIR = "/home/claude/pharma_project/outputs"
XLSX_PATH = "/home/claude/pharma_project/excel/Pharma_Analytics_Dashboard.xlsx"

NAVY, TEAL, GOLD, RED = "1F3864", "2E8B8B", "D4A017", "B23A48"
LIGHT, YELLOW, WHITE = "EDF1F7", "FFF2CC", "FFFFFF"
HEADER_FONT = Font(name="Arial", bold=True, color=WHITE, size=11)
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
TITLE_FONT = Font(name="Arial", bold=True, color=NAVY, size=18)
SUBTITLE_FONT = Font(name="Arial", italic=True, color="595959", size=10)
BODY_FONT = Font(name="Arial", size=10)
NOTE_FONT = Font(name="Arial", italic=True, size=9, color="808080")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
YELLOW_FILL = PatternFill("solid", fgColor=YELLOW)
RED_FILL = PatternFill("solid", fgColor="F8CBAD")

mol_perf = pd.read_csv(f"{OUT_DIR}/molecule_performance.csv")
monthly = pd.read_csv(f"{OUT_DIR}/monthly_trend.csv")
state_perf = pd.read_csv(f"{OUT_DIR}/state_performance.csv")
region_perf = pd.read_csv(f"{OUT_DIR}/region_performance.csv")
with open(f"{OUT_DIR}/underperforming_region.json") as f:
    underperf = json.load(f)

wb = load_workbook(XLSX_PATH)


def style_header_row(ws, row, ncols, start_col=1):
    for c in range(start_col, start_col + ncols):
        cell = ws.cell(row=row, column=c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER


def write_small_table(ws, df, start_row, start_col=1, table_name=None):
    for j, col in enumerate(df.columns):
        ws.cell(row=start_row, column=start_col + j, value=col)
    style_header_row(ws, start_row, len(df.columns), start_col)
    r0 = start_row + 1
    for i, row in enumerate(df.itertuples(index=False)):
        for j, val in enumerate(row):
            cell = ws.cell(row=r0 + i, column=start_col + j, value=val)
            cell.font = BODY_FONT
            cell.border = BORDER
    end_row = r0 + len(df) - 1
    if len(df) > 0:
        end_col_letter = get_column_letter(start_col + len(df.columns) - 1)
        start_col_letter = get_column_letter(start_col)
        ref = f"{start_col_letter}{start_row}:{end_col_letter}{end_row}"
        tbl = Table(displayName=table_name, ref=ref)
        tbl.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(tbl)
    return end_row


def autosize(ws, widths):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


# ====================================================================
# SHEET — Molecule and Therapy Uptake
# ====================================================================
ws = wb.create_sheet("Molecule and Therapy Uptake")
ws.sheet_view.showGridLines = False
ws["B2"] = "Molecule-Wise Performance & Therapy Uptake"
ws["B2"].font = TITLE_FONT
ws["B3"] = "SIMULATED DATA — growth rates precomputed in Python (first-6-month vs. last-6-month comparison)."
ws["B3"].font = SUBTITLE_FONT

r = 5
ws.cell(row=r, column=2, value="Molecule Performance & Growth").font = Font(name="Arial", bold=True, size=12, color=NAVY)
mol_display = mol_perf[["brand", "molecule", "therapy_area", "total_rx", "total_revenue", "growth_pct"]]
end1 = write_small_table(ws, mol_display, r + 1, start_col=2, table_name="MoleculePerf")
for rr in range(r + 2, end1 + 1):
    ws.cell(row=rr, column=5).number_format = '#,##0'
    ws.cell(row=rr, column=6).number_format = '$#,##0'
    growth_cell = ws.cell(row=rr, column=7)
    growth_cell.number_format = '+0.0"%";-0.0"%"'
    for cc in range(2, 8):
        ws.cell(row=rr, column=cc).fill = YELLOW_FILL
    if isinstance(growth_cell.value, (int, float)) and growth_cell.value < 0:
        growth_cell.fill = RED_FILL

r2 = end1 + 3
ws.cell(row=r2, column=2, value=(
    "Growth % = (avg. monthly prescriptions in the last 6 months) vs. (first 6 months), all "
    "doctors combined. Negative values (red) flag molecules losing share — e.g. Humira's simulated "
    "decline mirrors the real, well-documented 2023 biosimilar-driven erosion of Humira's market."
)).font = NOTE_FONT
ws.cell(row=r2, column=2).alignment = Alignment(wrap_text=True)
ws.row_dimensions[r2].height = 40

r3 = r2 + 3
ws.cell(row=r3, column=2, value="Overall Monthly Prescribing Trend").font = Font(name="Arial", bold=True, size=12, color=NAVY)
monthly_display = monthly.copy()
monthly_display["month"] = pd.to_datetime(monthly_display["month"]).dt.strftime("%Y-%m")
end2 = write_small_table(ws, monthly_display, r3 + 1, start_col=2, table_name="MonthlyTrend")
for rr in range(r3 + 2, end2 + 1):
    ws.cell(row=rr, column=3).number_format = '#,##0'
    ws.cell(row=rr, column=4).number_format = '$#,##0'
    for cc in range(2, 5):
        ws.cell(row=rr, column=cc).fill = YELLOW_FILL

autosize(ws, {"A": 2, "B": 16, "C": 22, "D": 22, "E": 13, "F": 14, "G": 11})

chart = LineChart()
chart.title = "Overall Monthly Prescribing Trend"
data_ref = Reference(ws, min_col=3, min_row=r3 + 1, max_row=end2)
cats_ref = Reference(ws, min_col=2, min_row=r3 + 2, max_row=end2)
chart.add_data(data_ref, titles_from_data=True)
chart.set_categories(cats_ref)
chart.height, chart.width = 9, 20
ws.add_chart(chart, "I5")

print("Molecule and Therapy Uptake sheet written.")

# ====================================================================
# SHEET — Regional Performance
# ====================================================================
ws = wb.create_sheet("Regional Performance")
ws.sheet_view.showGridLines = False
ws["B2"] = "Regional Performance & Underperforming-Region Flag"
ws["B2"].font = TITLE_FONT
ws["B3"] = "SIMULATED DATA — precomputed in Python (per-doctor normalization + portfolio benchmark)."
ws["B3"].font = SUBTITLE_FONT

r = 5
ws.cell(row=r, column=2, value=f"\u26a0 Flagged: {underperf['flagged_state']} \u2014 "
        f"{underperf['gap_pct_below_average']:.0f}% below portfolio average prescribing volume per doctor"
        ).font = Font(name="Arial", bold=True, size=12, color="9C4500")
for c in range(2, 8):
    ws.cell(row=r, column=c).fill = PatternFill("solid", fgColor="FCE4D6")
ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7)

r1 = r + 2
ws.cell(row=r1, column=2, value="State-Level Performance").font = Font(name="Arial", bold=True, size=12, color=NAVY)
state_display = state_perf[["state", "total_rx", "total_revenue", "doctors", "rx_per_doctor", "revenue_per_doctor"]]
end1 = write_small_table(ws, state_display, r1 + 1, start_col=2, table_name="StatePerf")
for rr in range(r1 + 2, end1 + 1):
    ws.cell(row=rr, column=3).number_format = '#,##0'
    ws.cell(row=rr, column=4).number_format = '$#,##0'
    ws.cell(row=rr, column=6).number_format = '#,##0'
    ws.cell(row=rr, column=7).number_format = '$#,##0'
    state_name = ws.cell(row=rr, column=2).value
    fill = RED_FILL if state_name == underperf["flagged_state"] else YELLOW_FILL
    for cc in range(2, 8):
        ws.cell(row=rr, column=cc).fill = fill

r2 = end1 + 3
ws.cell(row=r2, column=2, value="Opportunity Sizing (if flagged state reached portfolio average)").font = Font(
    name="Arial", bold=True, size=12, color=NAVY)
opp_rows = [
    ("Portfolio avg. prescriptions per doctor", underperf["portfolio_avg_rx_per_doctor"]),
    (f"{underperf['flagged_state']} prescriptions per doctor", underperf["state_rx_per_doctor"]),
    ("Gap below average", f"{underperf['gap_pct_below_average']:.1f}%"),
    ("Implied missed prescriptions (annualized gap)", underperf["implied_missed_rx_if_at_average"]),
    ("Implied missed revenue (simulated)", f"${underperf['implied_missed_revenue_if_at_average']:,.0f}"),
]
for i, (label, val) in enumerate(opp_rows):
    rr = r2 + 1 + i
    ws.cell(row=rr, column=2, value=label).font = BODY_FONT
    vcell = ws.cell(row=rr, column=3, value=val)
    vcell.font = Font(name="Arial", bold=True, size=10, color=NAVY)
    for cc in (2, 3):
        ws.cell(row=rr, column=cc).fill = YELLOW_FILL
        ws.cell(row=rr, column=cc).border = BORDER

autosize(ws, {"A": 2, "B": 34, "C": 16, "D": 16, "E": 10, "F": 16, "G": 16})

chart = BarChart()
chart.type = "bar"
chart.title = f"Prescriptions per Doctor by State (Portfolio avg: {underperf['portfolio_avg_rx_per_doctor']:.0f})"
data_ref = Reference(ws, min_col=6, min_row=r1 + 1, max_row=end1)
cats_ref = Reference(ws, min_col=2, min_row=r1 + 2, max_row=end1)
chart.add_data(data_ref, titles_from_data=True)
chart.set_categories(cats_ref)
chart.height, chart.width = 9, 18
ws.add_chart(chart, "I5")

wb.save(XLSX_PATH)
print("Regional Performance sheet written.")
