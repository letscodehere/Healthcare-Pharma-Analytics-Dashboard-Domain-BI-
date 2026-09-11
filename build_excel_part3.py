import pandas as pd
import json
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

OUT_DIR = "/home/claude/pharma_project/outputs"
XLSX_PATH = "/home/claude/pharma_project/excel/Pharma_Analytics_Dashboard.xlsx"

NAVY, TEAL, GOLD, RED = "1F3864", "2E8B8B", "D4A017", "B23A48"
YELLOW, WHITE = "FFF2CC", "FFFFFF"
HEADER_FONT = Font(name="Arial", bold=True, color=WHITE, size=11)
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
TITLE_FONT = Font(name="Arial", bold=True, color=NAVY, size=18)
SUBTITLE_FONT = Font(name="Arial", italic=True, color="595959", size=10)
BODY_FONT = Font(name="Arial", size=10)
NOTE_FONT = Font(name="Arial", italic=True, size=9, color="808080")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
YELLOW_FILL = PatternFill("solid", fgColor=YELLOW)

doc_perf = pd.read_csv(f"{OUT_DIR}/doctor_performance.csv")
budget = pd.read_csv(f"{OUT_DIR}/budget_reallocation.csv")
with open(f"{OUT_DIR}/budget_impact_projection.json") as f:
    impact = json.load(f)

wb = load_workbook(XLSX_PATH)


def style_header_row(ws, row, ncols, start_col=1):
    for c in range(start_col, start_col + ncols):
        cell = ws.cell(row=row, column=c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER


def write_small_table(ws, df, start_row, start_col=1, table_name=None):
    for j, col in enumerate(df.columns):
        ws.cell(row=start_row, column=start_col + j, value=col)
    style_header_row(ws, start_row, len(df.columns), start_col)
    r0 = start_row + 1
    for i, row in enumerate(df.itertuples(index=False)):
        for j, val in enumerate(row):
            cell = ws.cell(row=r0 + i, column=start_col + j, value=val)
            cell.font = BODY_FONT
            cell.border = BORDER
    end_row = r0 + len(df) - 1
    if len(df) > 0:
        end_col_letter = get_column_letter(start_col + len(df.columns) - 1)
        start_col_letter = get_column_letter(start_col)
        ref = f"{start_col_letter}{start_row}:{end_col_letter}{end_row}"
        tbl = Table(displayName=table_name, ref=ref)
        tbl.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(tbl)
    return end_row


def autosize(ws, widths):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def kpi_card(ws, row, col, label, value, number_format=None):
    lbl_cell = ws.cell(row=row, column=col, value=label)
    lbl_cell.font = Font(name="Arial", bold=True, color="595959", size=9)
    val_cell = ws.cell(row=row + 1, column=col, value=value)
    val_cell.font = Font(name="Arial", bold=True, color=NAVY, size=16)
    if number_format:
        val_cell.number_format = number_format
    for rr in (row, row + 1):
        c = ws.cell(row=rr, column=col)
        c.fill = YELLOW_FILL
        c.border = BORDER
    ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + 1)
    ws.merge_cells(start_row=row + 1, start_column=col, end_row=row + 1, end_column=col + 1)


# ====================================================================
# SHEET — Doctor Performance
# ====================================================================
ws = wb.create_sheet("Doctor Performance")
ws.sheet_view.showGridLines = False
ws["B2"] = "Doctor-Wise Performance"
ws["B2"].font = TITLE_FONT
ws["B3"] = "SIMULATED DATA — synthetic doctor names, not real physicians. Precomputed ranking in Python."
ws["B3"].font = SUBTITLE_FONT

r = 5
ws.cell(row=r, column=2, value="Top 25 Doctors by Prescribing Revenue").font = Font(
    name="Arial", bold=True, size=12, color=NAVY)
doc_display = doc_perf.head(25)[["doctor_name", "specialty", "state", "total_rx", "total_revenue"]]
end1 = write_small_table(ws, doc_display, r + 1, start_col=2, table_name="TopDoctors")
for rr in range(r + 2, end1 + 1):
    ws.cell(row=rr, column=5).number_format = '#,##0'
    ws.cell(row=rr, column=6).number_format = '$#,##0'
    for cc in range(2, 7):
        ws.cell(row=rr, column=cc).fill = YELLOW_FILL

autosize(ws, {"A": 2, "B": 22, "C": 30, "D": 16, "E": 13, "F": 14})

chart = BarChart()
chart.type = "bar"
chart.title = "Top 15 Doctors by Revenue"
data_ref = Reference(ws, min_col=6, min_row=r + 1, max_row=r + 16)
cats_ref = Reference(ws, min_col=2, min_row=r + 2, max_row=r + 16)
chart.add_data(data_ref, titles_from_data=True)
chart.set_categories(cats_ref)
chart.height, chart.width = 10, 18
ws.add_chart(chart, "H5")

print("Doctor Performance sheet written.")

# ====================================================================
# SHEET — Budget Reallocation
# ====================================================================
ws = wb.create_sheet("Budget Reallocation")
ws.sheet_view.showGridLines = False
ws["B2"] = "Marketing Budget Reallocation Proposal"
ws["B2"].font = TITLE_FONT
ws["B3"] = "ILLUSTRATIVE PLANNING ASSUMPTIONS — not a guaranteed outcome. See Read Me and business report."
ws["B3"].font = SUBTITLE_FONT

kpi_card(ws, 5, 2, "OHIO BUDGET (current)", impact["ohio_budget_before"], '$#,##0')
kpi_card(ws, 5, 4, "OHIO BUDGET (proposed)", impact["ohio_budget_after"], '$#,##0')
kpi_card(ws, 5, 6, "PROJECTED INCREMENTAL REVENUE / YEAR", impact["projected_incremental_revenue_per_year"], '$#,##0')

ws.cell(row=8, column=2, value=(
    f"Assumption: {impact['reallocation_pct_of_non_ohio_budget']:.0f}% of the marketing budget in every "
    f"other state is redirected to Ohio (the flagged underperforming region). If this raises Ohio's "
    f"prescribing volume {impact['gap_closure_assumption_pct']:.0f}% of the way to the portfolio average "
    f"(a stated, conservative planning assumption, not a measured result), the projected impact is "
    f"{impact['projected_incremental_rx_per_year']:,.0f} incremental prescriptions and "
    f"${impact['projected_incremental_revenue_per_year']:,.0f} in incremental annual revenue (simulated)."
)).font = NOTE_FONT
ws.cell(row=8, column=2).alignment = Alignment(wrap_text=True)
ws.row_dimensions[8].height = 60

r0 = 11
ws.cell(row=r0, column=2, value="Proposed Budget by State").font = Font(name="Arial", bold=True, size=12, color=NAVY)
budget_display = budget[["state", "current_budget", "proposed_budget", "budget_change", "budget_change_pct"]]
end1 = write_small_table(ws, budget_display, r0 + 1, start_col=2, table_name="BudgetProposal")
for rr in range(r0 + 2, end1 + 1):
    ws.cell(row=rr, column=3).number_format = '$#,##0'
    ws.cell(row=rr, column=4).number_format = '$#,##0'
    ws.cell(row=rr, column=5).number_format = '+$#,##0;-$#,##0'
    ws.cell(row=rr, column=6).number_format = '+0.0"%";-0.0"%"'
    for cc in range(2, 7):
        ws.cell(row=rr, column=cc).fill = YELLOW_FILL

autosize(ws, {"A": 2, "B": 16, "C": 15, "D": 15, "E": 13, "F": 13, "G": 20, "H": 20})

chart = BarChart()
chart.type = "col"
chart.title = "Current vs. Proposed Marketing Budget by State"
data_ref = Reference(ws, min_col=3, max_col=4, min_row=r0 + 1, max_row=end1)
cats_ref = Reference(ws, min_col=2, min_row=r0 + 2, max_row=end1)
chart.add_data(data_ref, titles_from_data=True)
chart.set_categories(cats_ref)
chart.height, chart.width = 10, 18
ws.add_chart(chart, "H11")

print("Budget Reallocation sheet written.")

# ====================================================================
# Final tab order + colors
# ====================================================================
order = ["Read Me", "Dashboard", "Molecule and Therapy Uptake", "Regional Performance",
         "Doctor Performance", "Budget Reallocation", "Data"]
wb._sheets = [wb[name] for name in order]
wb.active = wb.sheetnames.index("Dashboard")
for name in wb.sheetnames:
    wb[name].sheet_properties.tabColor = NAVY if name != "Data" else "A6A6A6"

wb.save(XLSX_PATH)
print("Final workbook saved. Sheets:", wb.sheetnames)
